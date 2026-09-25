"""Import aus Excel/CSV und Tabellen-Erfassung (NK-161, F-90).

Wer seine Nebenkosten bisher in Excel gerechnet hat, tippt nicht zweihundert
Rechnungen, fünfzehn Wohnungen und drei Jahre Zählerstände noch einmal ab.
Dieses Modul liest eine Tabelle (XLSX oder CSV), ordnet ihre Spalten den
Feldern der Anwendung zu, prüft jede Zeile und übernimmt alles in **einer**
Transaktion -- oder nichts.

Ablauf (``/api/import/...``):

  1. **Vorlage** herunterladen (XLSX mit Blatt „Daten“, „Beispiel“ und
     „Hinweise“, oder CSV mit Semikolon und UTF-8-BOM, wie Excel sie liest).
  2. **Lesen:** die Datei wird am Inhalt erkannt (``uploads.pruefe_tabelle``),
     die erste nicht leere Zeile ist der Kopf, höchstens ``MAX_ZEILEN``
     Zeilen. Die gelesenen Zeilen liegen eine Stunde im Arbeitsordner.
  3. **Zuordnen:** ein Vorschlag nach Spaltennamen, änderbar.
  4. **Prüfen (Probelauf):** jede Zeile wird gelesen, aufgelöst (Immobilie,
     Wohnung, Kostenart, Mieter, Zähler über ihre Namen) und probeweise
     angelegt -- danach zurückgerollt. Fehler je Zeile mit Excel-Zeilennummer.
  5. **Übernehmen:** derselbe Lauf, festgeschrieben nur ohne einen einzigen
     Fehler.

Die Tabellen-Erfassung der Rechnungen (``/api/import/rechnungen/tabelle``)
schickt ihre Zeilen direkt, ohne Datei -- derselbe Prüf- und Anlegeweg.

Datum: ``31.12.2025``, ``31.12.25``, ``2025-12-31`` oder eine Excel-Datumszelle.
Betrag: ``1.234,56``, ``1234,56``, ``1234.56``, mit oder ohne ``€``.
"""

from __future__ import annotations

import csv
import io
import json
import re
import secrets
import time
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

from sqlalchemy.exc import SQLAlchemyError

from geld import GeldFehler, runde

MAX_ZEILEN = 5000
ARBEIT = 'import-arbeit'
GUELTIG_S = 3600
ZAHLUNGSARTEN = ('Miete', 'Nebenkostenvorauszahlung', 'Kaution', 'Nebenkostenzahlung')
ZAHLUNGSART_ALIASE = {
    'vorauszahlung': 'Nebenkostenvorauszahlung',
    'nkvorauszahlung': 'Nebenkostenvorauszahlung',
    'abschlag': 'Nebenkostenvorauszahlung',
    'nachzahlung': 'Nebenkostenzahlung',
    'abrechnung': 'Nebenkostenzahlung',
    'kaltmiete': 'Miete',
}


# Alltagsnamen aus Excel-Tabellen -> Kostenart der Anwendung (§ 2 BetrKV).
# Greift nur, wenn es keinen gleichnamigen Eintrag gibt; der Bericht nennt,
# wie die Zeile gelesen wurde.
KOSTENART_ALIASE = {
    'muell': 'Straßenreinigung und Müllbeseitigung',
    'muellabfuhr': 'Straßenreinigung und Müllbeseitigung',
    'abfall': 'Straßenreinigung und Müllbeseitigung',
    'strassenreinigung': 'Straßenreinigung und Müllbeseitigung',
    'wasser': 'Wasserversorgung',
    'frischwasser': 'Wasserversorgung',
    'kaltwasser': 'Wasserversorgung',
    'abwasser': 'Entwässerung',
    'kanal': 'Entwässerung',
    'regenwasser': 'Niederschlagswasser',
    'allgemeinstrom': 'Beleuchtung (Allgemeinstrom)',
    'hausstrom': 'Beleuchtung (Allgemeinstrom)',
    'treppenhausbeleuchtung': 'Beleuchtung (Allgemeinstrom)',
    'versicherung': 'Sach- und Haftpflichtversicherung',
    'gebaeudeversicherung': 'Sach- und Haftpflichtversicherung',
    'hausmeister': 'Hauswart',
    'reinigung': 'Gebäudereinigung und Ungezieferbekämpfung',
    'treppenhausreinigung': 'Gebäudereinigung und Ungezieferbekämpfung',
    'schornsteinfeger': 'Schornsteinreinigung',
    'kabel': 'Antenne, Breitband oder Glasfaser',
    'kabelfernsehen': 'Antenne, Breitband oder Glasfaser',
    'fahrstuhl': 'Aufzug',
    'garten': 'Gartenpflege',
}


class ImportFehler(Exception):
    """Die Datei oder die Anfrage taugt nicht; Meldung für Menschen."""


class ZeilenFehler(Exception):
    """Eine Zeile taugt nicht; die Meldung nennt das Feld."""


@dataclass(frozen=True)
class Feld:
    schluessel: str
    titel: str
    pflicht: bool
    beispiel: str
    hinweis: str
    aliase: tuple = ()


@dataclass(frozen=True)
class Art:
    titel: str
    felder: tuple
    erklaerung: str
    beispiele: tuple = field(default_factory=tuple)


_IMMOBILIE = Feld('immobilie', 'Immobilie', True, 'Musterhaus Lindenstraße',
                  'Name des Hauses, wie er in der Anwendung steht.', ('haus', 'objekt', 'gebäude'))
_WOHNUNG = Feld('wohnung', 'Wohnung', True, 'EG links',
                'Name der Einheit in diesem Haus.', ('einheit', 'whg', 'lage'))

ARTEN = {
    'wohnungen': Art(
        'Immobilien und Wohnungen',
        (
            Feld('immobilie', 'Immobilie', True, 'Musterhaus Lindenstraße',
                 'Name des Hauses. Gibt es das Haus noch nicht, wird es angelegt.',
                 ('haus', 'objekt', 'gebäude')),
            _WOHNUNG,
            Feld('flaeche', 'Fläche (m²)', True, '62,5', 'Wohn- oder Nutzfläche in m².',
                 ('qm', 'm2', 'm²', 'wohnfläche', 'fläche')),
            Feld('nutzungsart', 'Nutzungsart', False, 'Wohnraum',
                 'Wohnraum oder Gewerbe; leer heißt Wohnraum.', ('nutzung',)),
            Feld('selbstversorger', 'Eigene Heizung', False, 'nein',
                 'ja, wenn die Einheit ihre Wärme selbst erzeugt (eigene Gastherme).',
                 ('selbstversorger',)),
            Feld('eigennutzung', 'Selbst bewohnt', False, 'nein',
                 'ja, wenn Sie die Einheit selbst nutzen.', ('eigennutzung',)),
        ),
        'Je Zeile eine Wohnung. Häuser, die es noch nicht gibt, entstehen mit.',
        (('Musterhaus Lindenstraße', 'EG links', '62,5', 'Wohnraum', 'nein', 'nein'),
         ('Musterhaus Lindenstraße', 'OG rechts', '71', 'Wohnraum', 'nein', 'nein')),
    ),
    'mieter': Art(
        'Mieter',
        (
            _IMMOBILIE, _WOHNUNG,
            Feld('name', 'Name', True, 'Anna Mieterin', 'Name des Mieters oder der Mietpartei.',
                 ('mieter', 'mietpartei')),
            Feld('einzug', 'Einzug', True, '01.03.2024', 'Tag des Einzugs (Mietbeginn).',
                 ('mietbeginn', 'einzugsdatum', 'von')),
            Feld('auszug', 'Auszug', False, '', 'Tag des Auszugs; leer, solange er wohnt.',
                 ('mietende', 'auszugsdatum', 'bis')),
            Feld('personen', 'Personen', False, '2',
                 'Wie viele Personen im Haushalt leben (ab Einzug); leer heißt 1.',
                 ('personenanzahl', 'haushalt', 'köpfe')),
        ),
        'Je Zeile ein Mietverhältnis. Haus und Wohnung müssen schon angelegt sein.',
        (('Musterhaus Lindenstraße', 'EG links', 'Anna Mieterin', '01.03.2024', '', '2'),),
    ),
    'rechnungen': Art(
        'Rechnungen',
        (
            _IMMOBILIE,
            Feld('kostenart', 'Kostenart', True, 'Grundsteuer',
                 'Eine Kostenart der Anwendung, z. B. Grundsteuer, Wasserversorgung, Hauswart. '
                 'Übliche Namen wie „Müllabfuhr“ oder „Allgemeinstrom“ werden erkannt.',
                 ('kategorie', 'art', 'position')),
            Feld('betrag', 'Betrag (€)', True, '1.234,56', 'Rechnungsbetrag brutto.',
                 ('betrag', 'summe', 'kosten', 'euro', 'brutto')),
            Feld('von', 'Zeitraum von', True, '01.01.2025', 'Erster Tag des Zeitraums der Rechnung.',
                 ('beginn', 'zeitraum von', 'start', 'von')),
            Feld('bis', 'Zeitraum bis', True, '31.12.2025', 'Letzter Tag des Zeitraums.',
                 ('ende', 'zeitraum bis', 'bis')),
            Feld('rechnungsdatum', 'Rechnungsdatum', False, '15.01.2026', 'Datum auf dem Beleg.',
                 ('belegdatum', 'datum')),
            Feld('rechnungsnummer', 'Rechnungsnummer', False, 'GS-2025-17', 'Nummer auf dem Beleg.',
                 ('nummer', 'rechnungsnr', 'belegnummer', 'nr')),
            Feld('anbieter', 'Anbieter', False, 'Stadtkasse',
                 'Wer die Rechnung gestellt hat. Unbekannte Anbieter werden angelegt.',
                 ('versorger', 'lieferant', 'firma')),
            Feld('wohnung', 'Nur für Wohnung', False, '',
                 'Leer, wenn die Rechnung das ganze Haus betrifft.', ('wohnung', 'einheit')),
            Feld('beschreibung', 'Beschreibung', False, '', 'Freitext.', ('notiz', 'bemerkung', 'text')),
        ),
        'Je Zeile eine Rechnung. Heizkosten mit Anlage und CO2-Angaben erfassen Sie im Dialog.',
        (('Musterhaus Lindenstraße', 'Grundsteuer', '1.234,56', '01.01.2025', '31.12.2025',
          '15.01.2025', 'GS-2025-17', 'Stadtkasse', '', ''),
         ('Musterhaus Lindenstraße', 'Wasserversorgung', '486,00', '01.01.2025', '31.12.2025',
          '', '', 'Stadtwerke', '', '')),
    ),
    'zahlungen': Art(
        'Zahlungen',
        (
            _IMMOBILIE, _WOHNUNG,
            Feld('mieter', 'Mieter', True, 'Anna Mieterin', 'Name des Mieters, wie in der Anwendung.',
                 ('name', 'mietpartei')),
            Feld('datum', 'Datum', True, '03.01.2025', 'Tag des Zahlungseingangs.',
                 ('zahlungsdatum', 'eingang', 'valuta', 'buchungstag')),
            Feld('betrag', 'Betrag (€)', True, '180,00', 'Gezahlter Betrag.', ('summe', 'euro')),
            Feld('art', 'Art', True, 'Nebenkostenvorauszahlung',
                 'Miete, Nebenkostenvorauszahlung, Kaution oder Nebenkostenzahlung.',
                 ('typ', 'zahlungsart', 'zweck')),
        ),
        'Je Zeile ein Zahlungseingang.',
        (('Musterhaus Lindenstraße', 'EG links', 'Anna Mieterin', '03.01.2025', '180,00',
          'Nebenkostenvorauszahlung'),),
    ),
    'zaehlerstaende': Art(
        'Zählerstände',
        (
            _IMMOBILIE,
            Feld('zaehlernummer', 'Zählernummer', True, 'WZ-10442',
                 'Nummer des Zählers, wie in der Anwendung angelegt.', ('zähler', 'zaehler', 'nummer')),
            Feld('datum', 'Ablesedatum', True, '31.12.2025', 'Tag der Ablesung.',
                 ('datum', 'stichtag', 'abgelesen am')),
            Feld('stand', 'Stand', True, '1284,5', 'Abgelesener Wert (bei Zweitarif: Hochtarif).',
                 ('zählerstand', 'wert', 'ht')),
            Feld('stand_nt', 'Stand Niedertarif', False, '', 'Nur bei Zweitarifzählern.',
                 ('nt', 'niedertarif')),
            Feld('zwischenablesung', 'Zwischenablesung', False, 'nein',
                 'ja, wenn der Stand zu einem Mieterwechsel abgelesen wurde.', ('mieterwechsel',)),
        ),
        'Je Zeile ein Zählerstand. Die Zähler müssen schon angelegt sein.',
        (('Musterhaus Lindenstraße', 'WZ-10442', '31.12.2025', '1284,5', '', 'nein'),),
    ),
}


# --- Werte lesen -----------------------------------------------------------------

def normiert(text) -> str:
    """Vergleichsform für Namen und Spaltenköpfe: klein, ohne Leer- und Sonderzeichen."""
    text = unicodedata.normalize('NFKC', str(text or '')).lower()
    text = text.replace('ä', 'ae').replace('ö', 'oe').replace('ü', 'ue').replace('ß', 'ss')
    return re.sub(r'[^a-z0-9²]', '', text)


def als_text(wert) -> str:
    if wert is None:
        return ''
    if isinstance(wert, float) and wert.is_integer():
        return str(int(wert))
    if isinstance(wert, datetime):
        return wert.date().isoformat()
    if isinstance(wert, date):
        return wert.isoformat()
    return str(wert).strip()


def als_datum(wert, titel: str) -> date:
    if isinstance(wert, datetime):
        return wert.date()
    if isinstance(wert, date):
        return wert
    if isinstance(wert, (int, float)) and not isinstance(wert, bool) and 20000 < wert < 80000:
        # Excel-Seriennummer (Zelle ohne Datumsformat).
        return date(1899, 12, 30) + timedelta(days=int(wert))
    text = als_text(wert)
    for muster in ('%d.%m.%Y', '%Y-%m-%d', '%d.%m.%y', '%d/%m/%Y'):
        try:
            return datetime.strptime(text, muster).date()
        except ValueError:
            continue
    raise ZeilenFehler(f'{titel}: „{text}“ ist kein Datum (z. B. 31.12.2025).')


def _zahlentext(text: str) -> str:
    text = text.replace('€', '').replace('EUR', '').replace(' ', '').replace(' ', '')
    text = text.replace("'", '')
    if ',' in text:
        return text.replace('.', '').replace(',', '.')
    if text.count('.') > 1:
        return text.replace('.', '')
    return text


def als_betrag(wert, titel: str) -> Decimal:
    if isinstance(wert, bool):
        raise ZeilenFehler(f'{titel}: hier steht ja/nein statt eines Betrags.')
    roh = repr(wert) if isinstance(wert, float) else _zahlentext(als_text(wert))
    try:
        return runde(roh)
    except GeldFehler:
        raise ZeilenFehler(f'{titel}: „{als_text(wert)}“ ist kein Betrag (z. B. 1.234,56).') from None


def als_zahl(wert, titel: str) -> float:
    if isinstance(wert, (int, float)) and not isinstance(wert, bool):
        return float(wert)
    try:
        zahl = float(_zahlentext(als_text(wert)))
    except ValueError:
        raise ZeilenFehler(f'{titel}: „{als_text(wert)}“ ist keine Zahl.') from None
    if zahl != zahl or zahl in (float('inf'), float('-inf')):
        raise ZeilenFehler(f'{titel}: „{als_text(wert)}“ ist keine Zahl.')
    return zahl


def als_wahrheit(wert, titel: str) -> bool:
    text = normiert(als_text(wert))
    if text in ('', 'nein', 'n', '0', 'falsch', 'false', 'no'):
        return False
    if text in ('ja', 'j', 'x', '1', 'wahr', 'true', 'yes', 'y'):
        return True
    raise ZeilenFehler(f'{titel}: bitte ja oder nein, hier stand „{als_text(wert)}“.')


# --- Dateien lesen ---------------------------------------------------------------

def _leer(zeile) -> bool:
    return not any(als_text(w) for w in zeile)


def tabelle_lesen(inhalt: bytes) -> tuple[list[str], list[list], int]:
    """Kopf, Zeilen und die Excel-Nummer der ersten Datenzeile.

    Die erste nicht leere Zeile ist der Kopf. Leere Zeilen dazwischen
    bleiben stehen (``verarbeiten`` überspringt sie), damit die Nummern in
    den Fehlermeldungen die Zeilen sind, die Excel zeigt.
    """
    import uploads
    art = uploads.pruefe_tabelle(inhalt)
    zeilen = _xlsx_zeilen(inhalt) if art == 'xlsx' else _csv_zeilen(inhalt)
    anfang = next((i for i, z in enumerate(zeilen) if not _leer(z)), None)
    if anfang is None:
        raise ImportFehler('Die Tabelle ist leer.')
    while zeilen and _leer(zeilen[-1]):
        zeilen.pop()
    kopf = [als_text(w) for w in zeilen[anfang]]
    daten = zeilen[anfang + 1:]
    if len(daten) > MAX_ZEILEN:
        raise ImportFehler(f'Die Tabelle hat {len(daten)} Zeilen; höchstens {MAX_ZEILEN} gehen '
                           'auf einmal. Bitte teilen Sie sie auf.')
    breite = len(kopf)
    return kopf, [list(z[:breite]) + [None] * (breite - len(z)) for z in daten], anfang + 2


def _xlsx_zeilen(inhalt: bytes) -> list[list]:
    import openpyxl
    try:
        mappe = openpyxl.load_workbook(io.BytesIO(inhalt), read_only=True, data_only=True)
    except Exception as fehler:  # noqa: BLE001 -- openpyxl wirft vieles; alles heißt „kaputt“
        raise ImportFehler('Die Excel-Datei lässt sich nicht öffnen. Speichern Sie sie in Excel '
                           'noch einmal als „Excel-Arbeitsmappe (.xlsx)“.') from fehler
    try:
        blatt = mappe.worksheets[0]
        zeilen = []
        for zeile in blatt.iter_rows(values_only=True):
            zeilen.append(list(zeile))
            if len(zeilen) > MAX_ZEILEN + 50:
                break
        return zeilen
    finally:
        mappe.close()


def _csv_zeilen(inhalt: bytes) -> list[list]:
    for kodierung in ('utf-8-sig', 'cp1252'):
        try:
            text = inhalt.decode(kodierung)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - cp1252 dekodiert fast alles
        raise ImportFehler('Die CSV-Datei ist weder UTF-8 noch Windows-1252.')
    erste = text.split('\n', 1)[0]
    trenner = max((';', ',', '\t'), key=erste.count)
    return [z for z in csv.reader(io.StringIO(text), delimiter=trenner)][:MAX_ZEILEN + 50]


# --- Vorlagen --------------------------------------------------------------------

def vorlage_csv(art: str) -> bytes:
    puffer = io.StringIO()
    schreiber = csv.writer(puffer, delimiter=';', lineterminator='\r\n')
    schreiber.writerow([f.titel for f in ARTEN[art].felder])
    return ('﻿' + puffer.getvalue()).encode('utf-8')


def vorlage_xlsx(art: str, kostenarten: list[str] | None = None) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    beschreibung = ARTEN[art]
    mappe = openpyxl.Workbook()
    daten = mappe.active
    daten.title = 'Daten'
    beispiel = mappe.create_sheet('Beispiel')
    hinweise = mappe.create_sheet('Hinweise')
    fett = Font(bold=True)
    pflichtfarbe = PatternFill('solid', fgColor='E0E7FF')
    for blatt in (daten, beispiel):
        blatt.append([f.titel for f in beschreibung.felder])
        for spalte, feld in enumerate(beschreibung.felder, 1):
            zelle = blatt.cell(row=1, column=spalte)
            zelle.font = fett
            if feld.pflicht:
                zelle.fill = pflichtfarbe
            blatt.column_dimensions[zelle.column_letter].width = max(14, len(feld.titel) + 4)
        blatt.freeze_panes = 'A2'
    for zeile in beschreibung.beispiele:
        beispiel.append(list(zeile))
    hinweise.append(['Spalte', 'Pflicht', 'Beispiel', 'Was hinein gehört'])
    for zelle in hinweise[1]:
        zelle.font = fett
    for feld in beschreibung.felder:
        hinweise.append([feld.titel, 'ja' if feld.pflicht else '', feld.beispiel, feld.hinweis])
    hinweise.append([])
    hinweise.append([beschreibung.erklaerung])
    hinweise.append(['Tragen Sie Ihre Daten ins Blatt „Daten“ ein. Gelesen wird nur das erste '
                     'Blatt. Pflichtspalten sind farbig hinterlegt.'])
    if kostenarten:
        hinweise.append([])
        hinweise.append(['Kostenarten dieser Installation'])
        hinweise[hinweise.max_row][0].font = fett
        for name in kostenarten:
            hinweise.append([name])
    for buchstabe, breite in zip('ABCD', (22, 8, 26, 80)):
        hinweise.column_dimensions[buchstabe].width = breite
    puffer = io.BytesIO()
    mappe.save(puffer)
    return puffer.getvalue()


# --- Zuordnung -------------------------------------------------------------------

def zuordnung_vorschlagen(art: str, kopf: list[str]) -> dict:
    """Feld -> Spaltennummer, nach Spaltennamen. Jede Spalte höchstens einmal."""
    spalten = {normiert(name): i for i, name in enumerate(kopf) if normiert(name)}
    vergeben = set()
    vorschlag = {}
    for feld in ARTEN[art].felder:
        kandidaten = [feld.titel, feld.schluessel, *feld.aliase]
        for kandidat in kandidaten:
            index = spalten.get(normiert(kandidat))
            if index is not None and index not in vergeben:
                vorschlag[feld.schluessel] = index
                vergeben.add(index)
                break
    return vorschlag


def _zuordnung_pruefen(art: str, zuordnung: dict, breite: int) -> dict:
    felder = {f.schluessel: f for f in ARTEN[art].felder}
    sauber = {}
    for schluessel, index in (zuordnung or {}).items():
        if schluessel not in felder or index in (None, ''):
            continue
        try:
            index = int(index)
        except (TypeError, ValueError):
            raise ImportFehler(f'Die Zuordnung für {felder[schluessel].titel} ist ungültig.') from None
        if not 0 <= index < breite:
            raise ImportFehler(f'Die Zuordnung für {felder[schluessel].titel} zeigt auf keine Spalte.')
        sauber[schluessel] = index
    fehlend = [f.titel for f in felder.values() if f.pflicht and f.schluessel not in sauber]
    if fehlend:
        raise ImportFehler('Diese Pflichtspalten sind keiner Spalte zugeordnet: '
                           + ', '.join(fehlend) + '.')
    return sauber


# --- Arbeitsordner ---------------------------------------------------------------

def _arbeitsordner(basis: Path) -> Path:
    ordner = basis / ARBEIT
    ordner.mkdir(parents=True, exist_ok=True)
    grenze = time.time() - GUELTIG_S
    for alt in ordner.glob('*.json'):
        try:
            if alt.stat().st_mtime < grenze:
                alt.unlink()
        except OSError:
            continue
    return ordner


def ablegen(basis: Path, art: str, kopf: list, zeilen: list, versatz: int = 2) -> str:
    kennung = secrets.token_hex(16)
    werte = [[als_text(w) if isinstance(w, (date, datetime)) else w for w in z] for z in zeilen]
    (_arbeitsordner(basis) / f'{kennung}.json').write_text(
        json.dumps({'art': art, 'kopf': kopf, 'zeilen': werte, 'versatz': versatz},
                   ensure_ascii=False),
        encoding='utf-8')
    return kennung


def abholen(basis: Path, art: str, kennung: str) -> tuple[list, list, int]:
    if not re.fullmatch(r'[0-9a-f]{32}', kennung or ''):
        raise ImportFehler('Die gelesene Tabelle ist abgelaufen. Bitte die Datei noch einmal wählen.')
    datei = _arbeitsordner(basis) / f'{kennung}.json'
    if not datei.is_file():
        raise ImportFehler('Die gelesene Tabelle ist abgelaufen. Bitte die Datei noch einmal wählen.')
    inhalt = json.loads(datei.read_text(encoding='utf-8'))
    if inhalt['art'] != art:
        raise ImportFehler('Die gelesene Tabelle gehört zu einer anderen Importart.')
    return inhalt['kopf'], inhalt['zeilen'], inhalt.get('versatz', 2)


def verwerfen(basis: Path, kennung: str) -> None:
    if re.fullmatch(r'[0-9a-f]{32}', kennung or ''):
        (_arbeitsordner(basis) / f'{kennung}.json').unlink(missing_ok=True)


# --- Zeilen verarbeiten ----------------------------------------------------------

class _Aufloeser:
    """Namen -> Datensätze, mit dem, was in diesem Lauf schon angelegt wurde."""

    def __init__(self):
        from models import CostCategory, Property, Provider
        self.haeuser = {normiert(p.name): p for p in Property.query.all()}
        self.kostenarten = {normiert(k.name): k for k in CostCategory.query.all()}
        self.anbieter = {normiert(a.name): a for a in Provider.query.all()}
        self.angelegt = {'immobilien': 0, 'wohnungen': 0, 'mieter': 0, 'rechnungen': 0,
                         'anbieter': 0, 'zahlungen': 0, 'zaehlerstaende': 0}

    def haus(self, name: str, anlegen: bool = False):
        from models import Property, db
        treffer = self.haeuser.get(normiert(name))
        if treffer is None and anlegen:
            treffer = Property(name=name[:100], is_standalone=False)
            db.session.add(treffer)
            db.session.flush()
            self.haeuser[normiert(name)] = treffer
            self.angelegt['immobilien'] += 1
        if treffer is None:
            raise ZeilenFehler(f'Immobilie: „{name}“ gibt es nicht. Legen Sie sie zuerst an '
                               '(oder importieren Sie zuerst die Wohnungen).')
        return treffer

    @staticmethod
    def wohnung(haus, name: str, pflicht: bool = True):
        from models import Apartment
        for wohnung in Apartment.query.filter_by(property_id=haus.id).all():
            if normiert(wohnung.name) == normiert(name):
                return wohnung
        if pflicht:
            raise ZeilenFehler(f'Wohnung: „{name}“ gibt es in „{haus.name}“ nicht.')
        return None

    def kostenart(self, name: str, hinweise: list | None = None):
        schluessel = normiert(name)
        treffer = self.kostenarten.get(schluessel)
        if treffer is None:
            ziel = KOSTENART_ALIASE.get(schluessel)
            treffer = self.kostenarten.get(normiert(ziel)) if ziel else None
            if treffer is None and len(schluessel) >= 4:
                teil = [k for n, k in self.kostenarten.items() if schluessel in n]
                treffer = teil[0] if len(teil) == 1 else None
            if treffer is not None and hinweise is not None:
                hinweise.append(f'Kostenart „{name}“ als „{treffer.name}“ übernommen.')
        if treffer is None:
            namen = sorted(k.name for k in self.kostenarten.values())
            bekannte = ', '.join(namen[:6]) + (' …' if len(namen) > 6 else '')
            raise ZeilenFehler(f'Kostenart: „{name}“ gibt es nicht. Die Liste steht in der Vorlage '
                               f'(Blatt „Hinweise“) und unter Einstellungen; z. B. {bekannte}')
        return treffer

    def anbieter_zu(self, name: str):
        from models import Provider, db
        if not name:
            return None
        treffer = self.anbieter.get(normiert(name))
        if treffer is None:
            treffer = Provider(name=name[:100])
            db.session.add(treffer)
            db.session.flush()
            self.anbieter[normiert(name)] = treffer
            self.angelegt['anbieter'] += 1
        return treffer


def _wert(zeile: list, zuordnung: dict, schluessel: str):
    index = zuordnung.get(schluessel)
    return None if index is None else zeile[index]


def _pflicht_text(zeile, zuordnung, feld: Feld, maxlaenge: int = 200) -> str:
    text = als_text(_wert(zeile, zuordnung, feld.schluessel))
    if feld.pflicht and not text:
        raise ZeilenFehler(f'{feld.titel}: fehlt.')
    if len(text) > maxlaenge:
        raise ZeilenFehler(f'{feld.titel}: zu lang ({len(text)} Zeichen, erlaubt {maxlaenge}).')
    return text


def _zeile_wohnungen(z, zu, felder, aufl, hinweise):
    from models import Apartment, db
    from nutzung import ARTEN as NUTZUNGEN
    haus_name = _pflicht_text(z, zu, felder['immobilie'], 100)
    name = _pflicht_text(z, zu, felder['wohnung'], 100)
    flaeche = als_zahl(_wert(z, zu, 'flaeche'), felder['flaeche'].titel)
    if flaeche <= 0:
        raise ZeilenFehler(f'{felder["flaeche"].titel}: muss größer als 0 sein.')
    roh_art = normiert(als_text(_wert(z, zu, 'nutzungsart')))
    nutzungsart = 'wohnen'
    if roh_art:
        passend = [schluessel for schluessel, text in NUTZUNGEN.items()
                   if roh_art in (normiert(schluessel), normiert(text))]
        if not passend:
            raise ZeilenFehler(f'Nutzungsart: bitte {" oder ".join(NUTZUNGEN.values())}.')
        nutzungsart = passend[0]
    haus = aufl.haus(haus_name, anlegen=True)
    if aufl.wohnung(haus, name, pflicht=False) is not None:
        raise ZeilenFehler(f'Wohnung: „{name}“ gibt es in „{haus.name}“ schon.')
    db.session.add(Apartment(
        property_id=haus.id, name=name, sqm=flaeche, nutzungsart=nutzungsart,
        selbstversorger=als_wahrheit(_wert(z, zu, 'selbstversorger'), felder['selbstversorger'].titel),
        eigennutzung=als_wahrheit(_wert(z, zu, 'eigennutzung'), felder['eigennutzung'].titel)))
    db.session.flush()
    aufl.angelegt['wohnungen'] += 1


def _zeile_mieter(z, zu, felder, aufl, hinweise):
    from models import Haushaltsgroesse, Tenant, db
    haus = aufl.haus(_pflicht_text(z, zu, felder['immobilie'], 100))
    wohnung = aufl.wohnung(haus, _pflicht_text(z, zu, felder['wohnung'], 100))
    name = _pflicht_text(z, zu, felder['name'], 200)
    einzug = als_datum(_wert(z, zu, 'einzug'), felder['einzug'].titel)
    auszug_roh = _wert(z, zu, 'auszug')
    auszug = als_datum(auszug_roh, felder['auszug'].titel) if als_text(auszug_roh) else None
    if auszug and auszug < einzug:
        raise ZeilenFehler('Auszug: liegt vor dem Einzug.')
    personen_roh = als_text(_wert(z, zu, 'personen'))
    personen = 1
    if personen_roh:
        personen = int(als_zahl(personen_roh, felder['personen'].titel))
        if personen < 1:
            raise ZeilenFehler('Personen: mindestens 1.')
    for anderer in Tenant.query.filter_by(apartment_id=wohnung.id).all():
        if normiert(anderer.name) == normiert(name) and anderer.move_in_date == einzug:
            raise ZeilenFehler(f'Name: „{name}“ mit diesem Einzug gibt es in dieser Wohnung schon.')
        ende = anderer.move_out_date or date.max
        if anderer.gesperrt_bis is None and anderer.move_in_date <= (auszug or date.max) \
                and einzug <= ende:
            hinweise.append(f'Überschneidet sich mit „{anderer.name}“ in derselben Wohnung.')
    mieter = Tenant(apartment_id=wohnung.id, name=name, move_in_date=einzug, move_out_date=auszug)
    db.session.add(mieter)
    db.session.flush()
    if personen_roh:
        db.session.add(Haushaltsgroesse(tenant_id=mieter.id, gueltig_ab=einzug,
                                        personenanzahl=personen))
    aufl.angelegt['mieter'] += 1


def _zeile_rechnungen(z, zu, felder, aufl, hinweise):
    from models import CostInvoice, db
    haus = aufl.haus(_pflicht_text(z, zu, felder['immobilie'], 100))
    kostenart = aufl.kostenart(_pflicht_text(z, zu, felder['kostenart'], 100), hinweise)
    betrag = als_betrag(_wert(z, zu, 'betrag'), felder['betrag'].titel)
    von = als_datum(_wert(z, zu, 'von'), felder['von'].titel)
    bis = als_datum(_wert(z, zu, 'bis'), felder['bis'].titel)
    if bis < von:
        raise ZeilenFehler('Zeitraum bis: liegt vor dem Beginn.')
    datum_roh = _wert(z, zu, 'rechnungsdatum')
    rechnungsdatum = als_datum(datum_roh, felder['rechnungsdatum'].titel) \
        if als_text(datum_roh) else None
    nummer = _pflicht_text(z, zu, felder['rechnungsnummer'], 100)
    wohnung_name = _pflicht_text(z, zu, felder['wohnung'], 100)
    wohnung = aufl.wohnung(haus, wohnung_name) if wohnung_name else None
    anbieter = aufl.anbieter_zu(_pflicht_text(z, zu, felder['anbieter'], 100))
    gleich = CostInvoice.query.filter_by(
        property_id=haus.id, category_id=kostenart.id, start_date=von, end_date=bis).all()
    if any(r.amount == betrag and (r.invoice_number or '') == nummer for r in gleich):
        hinweise.append('Eine gleiche Rechnung (Kostenart, Zeitraum, Betrag, Nummer) gibt es schon.')
    db.session.add(CostInvoice(
        property_id=haus.id, category_id=kostenart.id, amount=betrag,
        start_date=von, end_date=bis, rechnungsdatum=rechnungsdatum, invoice_number=nummer,
        provider_id=anbieter.id if anbieter else None,
        apartment_id=wohnung.id if wohnung else None,
        description=_pflicht_text(z, zu, felder['beschreibung'], 1000)))
    db.session.flush()
    aufl.angelegt['rechnungen'] += 1


def _zeile_zahlungen(z, zu, felder, aufl, hinweise):
    from models import Payment, Tenant, db
    haus = aufl.haus(_pflicht_text(z, zu, felder['immobilie'], 100))
    wohnung = aufl.wohnung(haus, _pflicht_text(z, zu, felder['wohnung'], 100))
    name = _pflicht_text(z, zu, felder['mieter'], 200)
    datum = als_datum(_wert(z, zu, 'datum'), felder['datum'].titel)
    betrag = als_betrag(_wert(z, zu, 'betrag'), felder['betrag'].titel)
    art_roh = normiert(_pflicht_text(z, zu, felder['art'], 100))
    art = next((a for a in ZAHLUNGSARTEN if normiert(a) == art_roh), None) \
        or ZAHLUNGSART_ALIASE.get(art_roh)
    if art is None:
        raise ZeilenFehler(f'Art: bitte {", ".join(ZAHLUNGSARTEN)}.')
    kandidaten = [t for t in Tenant.query.filter_by(apartment_id=wohnung.id).all()
                  if normiert(t.name) == normiert(name) and t.gesperrt_bis is None]
    if not kandidaten:
        raise ZeilenFehler(f'Mieter: „{name}“ wohnt nicht in „{wohnung.name}“.')
    passend = [t for t in kandidaten if t.move_in_date <= datum
               and (t.move_out_date is None or datum <= t.move_out_date + timedelta(days=62))]
    mieter = (passend or kandidaten)[0]
    if not passend:
        hinweise.append(f'Die Zahlung liegt außerhalb der Mietzeit von „{mieter.name}“.')
    db.session.add(Payment(tenant_id=mieter.id, amount=betrag, payment_date=datum, type=art))
    db.session.flush()
    aufl.angelegt['zahlungen'] += 1


def _zeile_zaehlerstaende(z, zu, felder, aufl, hinweise):
    from heizung import ABLESUNG, ZWISCHENABLESUNG
    from models import Meter, MeterReading, db
    haus = aufl.haus(_pflicht_text(z, zu, felder['immobilie'], 100))
    nummer = _pflicht_text(z, zu, felder['zaehlernummer'], 100)
    zaehler = [m for m in Meter.query.filter_by(property_id=haus.id).all()
               if normiert(m.meter_number) == normiert(nummer)]
    if not zaehler:
        raise ZeilenFehler(f'Zählernummer: „{nummer}“ gibt es in „{haus.name}“ nicht.')
    datum = als_datum(_wert(z, zu, 'datum'), felder['datum'].titel)
    stand = als_zahl(_wert(z, zu, 'stand'), felder['stand'].titel)
    nt_roh = _wert(z, zu, 'stand_nt')
    stand_nt = als_zahl(nt_roh, felder['stand_nt'].titel) if als_text(nt_roh) else None
    if stand_nt is not None and not zaehler[0].has_dual_tariff:
        raise ZeilenFehler('Stand Niedertarif: der Zähler hat nur einen Tarif.')
    if any(r.reading_date == datum for r in zaehler[0].readings):
        raise ZeilenFehler(f'Ablesedatum: für {datum:%d.%m.%Y} gibt es schon einen Stand.')
    zwischen = als_wahrheit(_wert(z, zu, 'zwischenablesung'), felder['zwischenablesung'].titel)
    db.session.add(MeterReading(meter_id=zaehler[0].id, reading_date=datum, value=stand,
                                value_nt=stand_nt,
                                ablesungsart=ZWISCHENABLESUNG if zwischen else ABLESUNG))
    db.session.flush()
    aufl.angelegt['zaehlerstaende'] += 1


_VERARBEITER = {
    'wohnungen': _zeile_wohnungen,
    'mieter': _zeile_mieter,
    'rechnungen': _zeile_rechnungen,
    'zahlungen': _zeile_zahlungen,
    'zaehlerstaende': _zeile_zaehlerstaende,
}


def verarbeiten(art: str, kopf: list, zeilen: list, zuordnung: dict, uebernehmen: bool,
                zeilenversatz: int = 2) -> dict:
    """Jede Zeile prüfen und anlegen; festschreiben nur ohne Fehler und nur auf Wunsch.

    ``zeilenversatz``: die Nummer der ersten Datenzeile, wie Excel sie zeigt
    (Kopf in Zeile 1 → 2). Die Tabellen-Erfassung zählt ab 1.
    """
    from models import db
    if art not in ARTEN:
        raise ImportFehler('Diese Importart gibt es nicht.')
    zuordnung = _zuordnung_pruefen(art, zuordnung, len(kopf))
    felder = {f.schluessel: f for f in ARTEN[art].felder}
    aufl = _Aufloeser()
    fehler, hinweise_alle, gelesen = [], [], 0
    for nummer, zeile in enumerate(zeilen, zeilenversatz):
        if _leer(zeile):
            continue
        gelesen += 1
        hinweise = []
        # Keine Savepoints: unter pysqlite beginnt ein SAVEPOINT vor dem
        # ersten INSERT die Transaktion selbst, und sein RELEASE schriebe
        # alles fest -- auch im Probelauf. Jede Zeile wird deshalb ganz
        # geprüft, bevor sie schreibt; ein Fehler bricht nur die Zeile ab.
        try:
            _VERARBEITER[art](zeile, zuordnung, felder, aufl, hinweise)
        except ZeilenFehler as grund:
            fehler.append({'zeile': nummer, 'meldung': str(grund)})
        except SQLAlchemyError as grund:
            # Unerwartet (die Prüfung oben deckt die Regeln der Datenbank
            # ab): die Sitzung ist danach unbrauchbar, der Lauf endet hier.
            db.session.rollback()
            fehler.append({'zeile': nummer, 'meldung': 'Die Zeile ließ sich nicht anlegen; '
                           f'weitere Zeilen wurden nicht geprüft ({type(grund).__name__}).'})
            break
        hinweise_alle += [{'zeile': nummer, 'meldung': h} for h in hinweise]
    uebernommen = uebernehmen and not fehler and gelesen > 0
    if uebernommen:
        db.session.commit()
    else:
        db.session.rollback()
    return {
        'art': art,
        'zeilen': gelesen,
        'fehlerfrei': gelesen - len({f['zeile'] for f in fehler}),
        'fehler': fehler,
        'hinweise': hinweise_alle,
        'angelegt': {k: v for k, v in aufl.angelegt.items() if v},
        'uebernommen': uebernommen,
    }


# --- Routen ----------------------------------------------------------------------

def init_tabellenimport(app):
    from flask import Response, jsonify, request

    import datenordner
    from validation import EingabeFehler

    def _basis() -> Path:
        try:
            return datenordner.datenordner()
        except datenordner.DatenordnerFehler:
            return Path(app.config.get('LOCAL_TEMP_PATH') or '.')

    def _art(art: str) -> str:
        if art not in ARTEN:
            raise EingabeFehler('Diese Importart gibt es nicht.', 'art')
        return art

    def _fehler(text):
        return jsonify({'error': text}), 400

    @app.route('/api/import/arten', methods=['GET'])
    def import_arten():
        return jsonify([{
            'art': schluessel, 'titel': art.titel, 'erklaerung': art.erklaerung,
            'felder': [{'schluessel': f.schluessel, 'titel': f.titel, 'pflicht': f.pflicht,
                        'beispiel': f.beispiel, 'hinweis': f.hinweis} for f in art.felder],
        } for schluessel, art in ARTEN.items()])

    @app.route('/api/import/vorlage/<art>.<endung>', methods=['GET'])
    def import_vorlage(art, endung):
        _art(art)
        if endung == 'xlsx':
            from models import CostCategory
            namen = [k.name for k in CostCategory.query.order_by(CostCategory.name).all()] \
                if art == 'rechnungen' else None
            inhalt = vorlage_xlsx(art, namen)
            typ = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        elif endung == 'csv':
            inhalt, typ = vorlage_csv(art), 'text/csv; charset=utf-8'
        else:
            return _fehler('Vorlagen gibt es als .xlsx und .csv.')
        return Response(inhalt, mimetype=typ, headers={
            'Content-Disposition': f'attachment; filename="NebenkostenFix-Vorlage-{art}.{endung}"',
            'X-Content-Type-Options': 'nosniff',
        })

    @app.route('/api/import/<art>/lesen', methods=['POST'])
    def import_lesen(art):
        _art(art)
        datei = request.files.get('datei')
        if datei is None or not datei.filename:
            raise EingabeFehler('Bitte wählen Sie eine Datei aus.', 'datei')
        try:
            kopf, zeilen, versatz = tabelle_lesen(datei.read())
        except ImportFehler as fehler:
            return _fehler(str(fehler))
        kennung = ablegen(_basis(), art, kopf, zeilen, versatz)
        return jsonify({
            'kennung': kennung, 'kopf': kopf, 'zeilen': sum(1 for z in zeilen if not _leer(z)),
            'zuordnung': zuordnung_vorschlagen(art, kopf),
            'beispiel': [[als_text(w) for w in z] for z in zeilen if not _leer(z)][:5],
        })

    def _lauf(art, uebernehmen):
        _art(art)
        daten = request.get_json(silent=True) or {}
        try:
            kopf, zeilen, versatz = abholen(_basis(), art, daten.get('kennung'))
            bericht = verarbeiten(art, kopf, zeilen, daten.get('zuordnung') or {}, uebernehmen,
                                  zeilenversatz=versatz)
        except ImportFehler as fehler:
            return _fehler(str(fehler))
        if bericht['uebernommen']:
            verwerfen(_basis(), daten.get('kennung'))
        return jsonify(bericht), (400 if uebernehmen and not bericht['uebernommen'] else 200)

    @app.route('/api/import/<art>/pruefen', methods=['POST'])
    def import_pruefen(art):
        """Probelauf: alles prüfen und probeweise anlegen, dann zurückrollen."""
        return _lauf(art, False)

    @app.route('/api/import/<art>/uebernehmen', methods=['POST'])
    def import_uebernehmen(art):
        """In einer Transaktion festschreiben -- nur ohne einen einzigen Fehler."""
        return _lauf(art, True)

    @app.route('/api/import/<art>/tabelle', methods=['POST'])
    def import_tabelle(art):
        """Tabellen-Erfassung (Zeilen direkt, ohne Datei): Zeilen als Objekte
        mit den Feldschlüsseln der Art; ``uebernehmen`` wie oben."""
        _art(art)
        daten = request.get_json(silent=True) or {}
        roh = daten.get('zeilen')
        if not isinstance(roh, list) or not roh:
            return _fehler('Die Tabelle ist leer.')
        if len(roh) > MAX_ZEILEN:
            return _fehler(f'Höchstens {MAX_ZEILEN} Zeilen auf einmal.')
        kopf = [f.schluessel for f in ARTEN[art].felder]
        zeilen = [[(z.get(k) if isinstance(z, dict) else None) for k in kopf] for z in roh]
        try:
            bericht = verarbeiten(art, kopf, zeilen, {k: i for i, k in enumerate(kopf)},
                                  bool(daten.get('uebernehmen')), zeilenversatz=1)
        except ImportFehler as fehler:
            return _fehler(str(fehler))
        return jsonify(bericht), (400 if daten.get('uebernehmen') and not bericht['uebernommen']
                                  else 200)
