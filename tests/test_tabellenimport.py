"""NK-161: Import aus Excel/CSV und Tabellen-Erfassung.

Die Werte, wie sie in deutschen Tabellen stehen (1.234,56 · 31.12.2025 ·
Excel-Datumszellen), jede Importart mit Anlegen und Probelauf, alles oder
nichts, Fehler mit Excel-Zeilennummer -- und die Dateien, die abgelehnt
gehören (altes .xls, .ods, PDF, ZIP-Bombe, zu viele Zeilen).
"""

from __future__ import annotations

import io
import zipfile
from datetime import date, datetime
from decimal import Decimal

import pytest

from nebenkostenfix import tabellenimport as ti


# --- Werte -----------------------------------------------------------------------

@pytest.mark.parametrize('roh, erwartet', [
    ('31.12.2025', date(2025, 12, 31)), ('1.3.2024', date(2024, 3, 1)),
    ('2025-12-31', date(2025, 12, 31)), ('31.12.25', date(2025, 12, 31)),
    (datetime(2025, 1, 2, 0, 0), date(2025, 1, 2)), (45658, date(2025, 1, 1)),
])
def test_datum(roh, erwartet):
    assert ti.als_datum(roh, 'Datum') == erwartet


@pytest.mark.parametrize('roh', ['32.01.2025', 'morgen', '12', ''])
def test_kein_datum(roh):
    with pytest.raises(ti.ZeilenFehler, match='kein Datum'):
        ti.als_datum(roh, 'Datum')


@pytest.mark.parametrize('roh, erwartet', [
    ('1.234,56', '1234.56'), ('1234,56 €', '1234.56'), ('1234.56', '1234.56'),
    ('1.234.567', '1234567.00'), (1234.5, '1234.50'), (486, '486.00'), ('€ 12', '12.00'),
    ('1 234,5', '1234.50'),
])
def test_betrag(roh, erwartet):
    assert ti.als_betrag(roh, 'Betrag') == Decimal(erwartet)


def test_kein_betrag():
    with pytest.raises(ti.ZeilenFehler, match='kein Betrag'):
        ti.als_betrag('zwölf', 'Betrag')
    with pytest.raises(ti.ZeilenFehler, match='ja/nein'):
        ti.als_betrag(True, 'Betrag')


def test_wahrheit():
    assert ti.als_wahrheit('Ja', 'x') is True and ti.als_wahrheit('x', 'x') is True
    assert ti.als_wahrheit('', 'x') is False and ti.als_wahrheit('nein', 'x') is False
    with pytest.raises(ti.ZeilenFehler, match='ja oder nein'):
        ti.als_wahrheit('vielleicht', 'x')


def test_zuordnung_nach_spaltennamen():
    kopf = ['Objekt', 'Whg', 'qm', 'Bemerkung']
    assert ti.zuordnung_vorschlagen('wohnungen', kopf) == {'immobilie': 0, 'wohnung': 1,
                                                           'flaeche': 2}


def test_openpyxl_liest_mit_defusedxml():
    import openpyxl.xml
    assert openpyxl.xml.DEFUSEDXML is True


# --- Vorlagen --------------------------------------------------------------------

@pytest.mark.parametrize('art', sorted(ti.ARTEN))
def test_vorlagen_lassen_sich_wieder_lesen(auth_client, art):
    import openpyxl
    antwort = auth_client.get(f'/api/import/vorlage/{art}.xlsx')
    assert antwort.status_code == 200
    assert 'NebenkostenFix-Vorlage' in antwort.headers['Content-Disposition']
    mappe = openpyxl.load_workbook(io.BytesIO(antwort.data))
    assert mappe.sheetnames == ['Daten', 'Beispiel', 'Hinweise']
    titel = [f.titel for f in ti.ARTEN[art].felder]
    assert [z.value for z in mappe['Daten'][1]] == titel
    assert mappe['Daten'].max_row == 1  # kein Beispiel im Datenblatt
    kopf, zeilen, versatz = ti.tabelle_lesen(antwort.data)
    assert kopf == titel and zeilen == [] and versatz == 2
    assert ti.zuordnung_vorschlagen(art, kopf) == {f.schluessel: i for i, f in
                                                    enumerate(ti.ARTEN[art].felder)}

    spalte_a = [z.value for z in mappe['Hinweise']['A']]
    if art == 'rechnungen':
        assert 'Kostenarten dieser Installation' in spalte_a and 'Grundsteuer' in spalte_a
    csv = auth_client.get(f'/api/import/vorlage/{art}.csv')
    assert csv.data.startswith('﻿'.encode()) and b';' in csv.data
    assert auth_client.get(f'/api/import/vorlage/{art}.ods').status_code == 400


def test_unbekannte_art(auth_client):
    assert auth_client.get('/api/import/vorlage/konten.xlsx').status_code == 400


# --- Durchlauf über die API ------------------------------------------------------

def _csv(*zeilen: str) -> bytes:
    return ('﻿' + '\r\n'.join(zeilen) + '\r\n').encode('utf-8')


def _lesen(client, art, inhalt, name='daten.csv'):
    antwort = client.post(f'/api/import/{art}/lesen',
                          data={'datei': (io.BytesIO(inhalt), name)},
                          content_type='multipart/form-data')
    assert antwort.status_code == 200, antwort.get_data(as_text=True)
    return antwort.get_json()


def _lauf(client, art, gelesen, schritt='pruefen', zuordnung=None):
    return client.post(f'/api/import/{art}/{schritt}', json={
        'kennung': gelesen['kennung'], 'zuordnung': zuordnung or gelesen['zuordnung']})


WOHNUNGEN = _csv('Immobilie;Wohnung;Fläche (m²);Nutzungsart;Eigene Heizung;Selbst bewohnt',
                 'Musterhaus Lindenstraße;EG links;62,5;Wohnraum;nein;nein',
                 'Musterhaus Lindenstraße;OG rechts;71;;;ja',
                 ';;;;;',
                 'Gartenhaus;Laden;40;Gewerbe;ja;')


def test_wohnungen_probelauf_aendert_nichts_uebernahme_legt_an(auth_client):
    from nebenkostenfix.models import Apartment, Property
    gelesen = _lesen(auth_client, 'wohnungen', WOHNUNGEN)
    assert gelesen['zeilen'] == 3 and len(gelesen['zuordnung']) == 6

    probe = _lauf(auth_client, 'wohnungen', gelesen).get_json()
    assert probe['fehler'] == [] and probe['zeilen'] == 3 and probe['uebernommen'] is False
    assert probe['angelegt'] == {'immobilien': 2, 'wohnungen': 3}
    assert Property.query.count() == 0 and Apartment.query.count() == 0

    antwort = _lauf(auth_client, 'wohnungen', gelesen, 'uebernehmen')
    assert antwort.status_code == 200 and antwort.get_json()['uebernommen'] is True
    wohnungen = {a.name: a for a in Apartment.query.all()}
    assert wohnungen['EG links'].sqm == 62.5 and wohnungen['OG rechts'].eigennutzung is True
    assert wohnungen['Laden'].nutzungsart == 'gewerbe' and wohnungen['Laden'].selbstversorger
    # Die gelesene Tabelle ist verbraucht.
    assert _lauf(auth_client, 'wohnungen', gelesen, 'uebernehmen').status_code == 400

    # Noch einmal: jede Wohnung gibt es jetzt schon.
    wieder = _lauf(auth_client, 'wohnungen', _lesen(auth_client, 'wohnungen', WOHNUNGEN)).get_json()
    assert [f['zeile'] for f in wieder['fehler']] == [2, 3, 5]
    assert 'schon' in wieder['fehler'][0]['meldung']


def test_alles_oder_nichts_mit_zeilennummern(auth_client):
    from nebenkostenfix.models import Apartment
    inhalt = _csv('Immobilie;Wohnung;Fläche (m²)',
                  'Musterhaus Lindenstraße;EG links;62,5',
                  'Musterhaus Lindenstraße;OG;viel',
                  'Musterhaus Lindenstraße;;50',
                  'Musterhaus Lindenstraße;DG;-3')
    gelesen = _lesen(auth_client, 'wohnungen', inhalt)
    antwort = _lauf(auth_client, 'wohnungen', gelesen, 'uebernehmen')
    assert antwort.status_code == 400
    bericht = antwort.get_json()
    assert bericht['uebernommen'] is False and bericht['fehlerfrei'] == 1
    assert [(f['zeile'], f['meldung'].split(':')[0]) for f in bericht['fehler']] == [
        (3, 'Fläche (m²)'), (4, 'Wohnung'), (5, 'Fläche (m²)')]
    assert Apartment.query.count() == 0


def test_pflichtspalte_ohne_zuordnung(auth_client):
    gelesen = _lesen(auth_client, 'wohnungen', _csv('Haus;Wohnung', 'A;B'))
    antwort = _lauf(auth_client, 'wohnungen', gelesen)
    assert antwort.status_code == 400 and 'Fläche (m²)' in antwort.get_json()['error']
    falsch = _lauf(auth_client, 'wohnungen', gelesen,
                   zuordnung={'immobilie': 0, 'wohnung': 1, 'flaeche': 9})
    assert 'keine Spalte' in falsch.get_json()['error']


@pytest.fixture
def haus(auth_client):
    from nebenkostenfix.models import Apartment, Meter, Property, db
    from nebenkostenfix.models import CostCategory
    haus = Property(name='Musterhaus Lindenstraße')
    db.session.add(haus)
    db.session.flush()
    eg = Apartment(property_id=haus.id, name='EG links', sqm=62.5)
    db.session.add(eg)
    wasser = CostCategory.query.filter(CostCategory.name.ilike('%wasser%')).first()
    db.session.add(Meter(property_id=haus.id, category_id=wasser.id, meter_number='WZ-10442',
                         is_main_meter=True))
    db.session.commit()
    return auth_client


def test_mieter_mit_personen_und_fehlern(haus):
    from nebenkostenfix.models import Tenant
    inhalt = _csv('Immobilie;Wohnung;Name;Einzug;Auszug;Personen',
                  'Musterhaus Lindenstraße;EG links;Anna Mieterin;01.03.2024;;2',
                  'Musterhaus Lindenstraße;EG rechts;Ben Beispiel;01.03.2024;;',
                  'Musterhaus Lindenstraße;EG links;Carla;01.05.2024;01.04.2024;',
                  'Nirgendhaus;EG links;Dora;01.01.2024;;')
    gelesen = _lesen(haus, 'mieter', inhalt)
    bericht = _lauf(haus, 'mieter', gelesen).get_json()
    meldungen = {f['zeile']: f['meldung'] for f in bericht['fehler']}
    assert 'EG rechts' in meldungen[3] and 'vor dem Einzug' in meldungen[4]
    assert 'Nirgendhaus' in meldungen[5]

    ok = _lesen(haus, 'mieter', _csv(
        'Immobilie;Wohnung;Name;Einzug;Auszug;Personen',
        'Musterhaus Lindenstraße;EG links;Anna Mieterin;01.03.2024;;2'))
    assert _lauf(haus, 'mieter', ok, 'uebernehmen').status_code == 200
    anna = Tenant.query.filter_by(name='Anna Mieterin').one()
    assert anna.move_in_date == date(2024, 3, 1) and anna.personenanzahl_am(date(2024, 6, 1)) == 2


def test_rechnungen_aus_excel_mit_datums_und_zahlenzellen(haus):
    import openpyxl
    from nebenkostenfix.models import CostInvoice, Provider
    mappe = openpyxl.Workbook()
    blatt = mappe.active
    blatt.append(['Objekt', 'Kategorie', 'Summe', 'Beginn', 'Ende', 'Nr', 'Firma'])
    blatt.append(['Musterhaus Lindenstraße', 'Grundsteuer', 1234.56, datetime(2025, 1, 1),
                  datetime(2025, 12, 31), 'GS-17', 'Stadtkasse'])
    blatt.append(['Musterhaus Lindenstraße', 'grundsteuer', '486,00', '01.01.2025',
                  '31.12.2025', None, None])
    blatt.append(['Musterhaus Lindenstraße', 'Kaviar', '1', '01.01.2025', '31.12.2025', None, None])
    puffer = io.BytesIO()
    mappe.save(puffer)
    gelesen = _lesen(haus, 'rechnungen', puffer.getvalue(), 'rechnungen.xlsx')
    assert gelesen['zuordnung'] == {'immobilie': 0, 'kostenart': 1, 'betrag': 2, 'von': 3,
                                    'bis': 4, 'rechnungsnummer': 5, 'anbieter': 6}
    bericht = _lauf(haus, 'rechnungen', gelesen).get_json()
    assert [f['zeile'] for f in bericht['fehler']] == [4]
    assert 'Kaviar' in bericht['fehler'][0]['meldung']
    assert 'Hinweise' in bericht['fehler'][0]['meldung']  # sagt, wo die Liste steht
    assert 'Aufzug' in bericht['fehler'][0]['meldung']
    assert Provider.query.count() == 0

    ok = _lesen(haus, 'rechnungen', puffer.getvalue(), 'rechnungen.xlsx')
    zuordnung = dict(ok['zuordnung'])
    ok_zeilen = _lauf(haus, 'rechnungen', ok, 'uebernehmen', zuordnung)
    assert ok_zeilen.status_code == 400  # die Kaviar-Zeile hält alles auf
    assert CostInvoice.query.count() == 0


def test_rechnungen_uebernehmen_und_doppelte_melden(haus):
    from nebenkostenfix.models import CostInvoice, Provider
    inhalt = _csv('Immobilie;Kostenart;Betrag (€);Zeitraum von;Zeitraum bis;Rechnungsnummer;Anbieter;Nur für Wohnung',
                  'Musterhaus Lindenstraße;Grundsteuer;1.234,56;01.01.2025;31.12.2025;GS-17;Stadtkasse;',
                  'Musterhaus Lindenstraße;Müllabfuhr;486;01.01.2025;31.12.2025;;Stadtkasse;EG links')
    gelesen = _lesen(haus, 'rechnungen', inhalt)
    bericht = _lauf(haus, 'rechnungen', gelesen, 'uebernehmen').get_json()
    assert bericht['fehler'] == [], bericht
    assert bericht['uebernommen'] and bericht['angelegt'] == {'rechnungen': 2, 'anbieter': 1}
    assert Provider.query.one().name == 'Stadtkasse'
    rechnungen = CostInvoice.query.order_by(CostInvoice.id).all()
    assert rechnungen[0].amount == Decimal('1234.56') and rechnungen[1].apartment_id is not None
    assert rechnungen[1].category.name == 'Straßenreinigung und Müllbeseitigung'
    assert 'als „Straßenreinigung und Müllbeseitigung“' in bericht['hinweise'][0]['meldung']
    wieder = _lauf(haus, 'rechnungen', _lesen(haus, 'rechnungen', inhalt)).get_json()
    doppelt = [h['zeile'] for h in wieder['hinweise'] if 'gleiche Rechnung' in h['meldung']]
    assert wieder['fehler'] == [] and doppelt == [2, 3]


def test_zahlungen_mit_kurzform_der_art(haus):
    from nebenkostenfix.models import Apartment, Payment, Tenant, db
    eg = Apartment.query.filter_by(name='EG links').one()
    db.session.add(Tenant(apartment_id=eg.id, name='Anna Mieterin', move_in_date=date(2024, 1, 1)))
    db.session.commit()
    inhalt = _csv('Immobilie;Wohnung;Mieter;Datum;Betrag (€);Art',
                  'Musterhaus Lindenstraße;EG links;Anna Mieterin;03.01.2025;180,00;Vorauszahlung',
                  'Musterhaus Lindenstraße;EG links;Anna Mieterin;03.02.2025;180,00;Nebenkostenvorauszahlung',
                  'Musterhaus Lindenstraße;EG links;Ben;03.02.2025;180,00;Miete',
                  'Musterhaus Lindenstraße;EG links;Anna Mieterin;03.02.2025;180,00;Taschengeld')
    bericht = _lauf(haus, 'zahlungen', _lesen(haus, 'zahlungen', inhalt)).get_json()
    assert [f['zeile'] for f in bericht['fehler']] == [4, 5]
    ok = _csv('Immobilie;Wohnung;Mieter;Datum;Betrag (€);Art',
              'Musterhaus Lindenstraße;EG links;anna mieterin;03.01.2025;180,00;Vorauszahlung')
    assert _lauf(haus, 'zahlungen', _lesen(haus, 'zahlungen', ok), 'uebernehmen').status_code == 200
    zahlung = Payment.query.one()
    assert zahlung.type == 'Nebenkostenvorauszahlung' and zahlung.amount == Decimal('180.00')


def test_zaehlerstaende(haus):
    from nebenkostenfix.models import MeterReading
    inhalt = _csv('Immobilie;Zählernummer;Ablesedatum;Stand;Stand Niedertarif;Zwischenablesung',
                  'Musterhaus Lindenstraße;WZ-10442;31.12.2025;1284,5;;nein',
                  'Musterhaus Lindenstraße;WZ-10442;30.06.2025;1100;;ja',
                  'Musterhaus Lindenstraße;WZ-10442;31.03.2025;1000;400;',
                  'Musterhaus Lindenstraße;XX-1;31.12.2025;1;;')
    bericht = _lauf(haus, 'zaehlerstaende', _lesen(haus, 'zaehlerstaende', inhalt)).get_json()
    assert {f['zeile']: f['meldung'].split(':')[0] for f in bericht['fehler']} == {
        4: 'Stand Niedertarif', 5: 'Zählernummer'}
    ok = _csv('Immobilie;Zählernummer;Ablesedatum;Stand;Stand Niedertarif;Zwischenablesung',
              'Musterhaus Lindenstraße;WZ-10442;31.12.2025;1284,5;;nein',
              'Musterhaus Lindenstraße;WZ-10442;30.06.2025;1100;;ja')
    assert _lauf(haus, 'zaehlerstaende', _lesen(haus, 'zaehlerstaende', ok),
                 'uebernehmen').status_code == 200
    arten = {r.reading_date: r.ablesungsart for r in MeterReading.query.all()}
    assert arten == {date(2025, 12, 31): 'ablesung', date(2025, 6, 30): 'zwischenablesung'}


def test_tabellen_erfassung_der_rechnungen(haus):
    from nebenkostenfix.models import CostInvoice
    zeilen = [
        {'immobilie': 'Musterhaus Lindenstraße', 'kostenart': 'Grundsteuer', 'betrag': '1.234,56',
         'von': '2025-01-01', 'bis': '2025-12-31'},
        {},
        {'immobilie': 'Musterhaus Lindenstraße', 'kostenart': 'Müllabfuhr', 'betrag': 'x',
         'von': '2025-01-01', 'bis': '2025-12-31'},
    ]
    probe = haus.post('/api/import/rechnungen/tabelle', json={'zeilen': zeilen}).get_json()
    assert probe['zeilen'] == 2 and [f['zeile'] for f in probe['fehler']] == [3]
    antwort = haus.post('/api/import/rechnungen/tabelle', json={'zeilen': zeilen[:1],
                                                                'uebernehmen': True})
    assert antwort.status_code == 200 and CostInvoice.query.count() == 1
    assert haus.post('/api/import/rechnungen/tabelle', json={'zeilen': []}).status_code == 400


# --- Ablehnungen -----------------------------------------------------------------

def test_abgelaufene_kennung(auth_client):
    antwort = auth_client.post('/api/import/wohnungen/pruefen',
                               json={'kennung': 'a' * 32, 'zuordnung': {}})
    assert antwort.status_code == 400 and 'abgelaufen' in antwort.get_json()['error']
    assert auth_client.post('/api/import/wohnungen/pruefen',
                            json={'kennung': '../x'}).status_code == 400


def test_kennung_gilt_nur_fuer_ihre_art(auth_client):
    gelesen = _lesen(auth_client, 'wohnungen', WOHNUNGEN)
    antwort = _lauf(auth_client, 'mieter', gelesen)
    assert antwort.status_code == 400 and 'anderen Importart' in antwort.get_json()['error']


@pytest.mark.parametrize('inhalt, muster', [
    (b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'\0' * 100, 'alte Excel-Format'),
    (b'%PDF-1.4\n...', 'keine Tabelle'),
    (b'a;b\n\x00\x00', 'keine Tabelle'),
    (b'   ', 'keine Tabelle'),
])
def test_falsche_dateien(auth_client, inhalt, muster):
    antwort = auth_client.post('/api/import/wohnungen/lesen',
                               data={'datei': (io.BytesIO(inhalt), 'x.xlsx')},
                               content_type='multipart/form-data')
    assert antwort.status_code == 400 and muster in antwort.get_json()['error']


def _zip(eintraege: dict) -> bytes:
    puffer = io.BytesIO()
    with zipfile.ZipFile(puffer, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, daten in eintraege.items():
            z.writestr(name, daten)
    return puffer.getvalue()


def test_ods_und_fremde_zip_und_bombe(auth_client):
    from nebenkostenfix import uploads
    with pytest.raises(Exception, match='OpenDocument'):
        uploads.pruefe_tabelle(_zip({'mimetype': 'x', 'content.xml': '<x/>'}))
    with pytest.raises(Exception, match='keine Tabelle'):
        uploads.pruefe_tabelle(_zip({'foto.jpg': 'x'}))
    bombe = _zip({'[Content_Types].xml': '<x/>', 'xl/workbook.xml': '<x/>',
                  'xl/worksheets/sheet1.xml': '\0' * (uploads.TABELLE_MAX_ENTPACKT + 1)})
    assert len(bombe) < 200_000
    with pytest.raises(Exception, match='zu groß'):
        uploads.pruefe_tabelle(bombe)


def test_kaputte_xlsx(auth_client):
    kaputt = _zip({'[Content_Types].xml': 'kein xml', 'xl/workbook.xml': 'auch nicht'})
    antwort = auth_client.post('/api/import/wohnungen/lesen',
                               data={'datei': (io.BytesIO(kaputt), 'x.xlsx')},
                               content_type='multipart/form-data')
    assert antwort.status_code == 400 and 'nicht öffnen' in antwort.get_json()['error']


def test_zu_viele_zeilen(auth_client, monkeypatch):
    monkeypatch.setattr(ti, 'MAX_ZEILEN', 3)
    inhalt = _csv('Immobilie;Wohnung;Fläche (m²)', *['H;W{};1'.format(i) for i in range(5)])
    antwort = auth_client.post('/api/import/wohnungen/lesen',
                               data={'datei': (io.BytesIO(inhalt), 'x.csv')},
                               content_type='multipart/form-data')
    assert antwort.status_code == 400 and 'höchstens 3' in antwort.get_json()['error']


def test_csv_in_windows_1252_mit_komma(auth_client):
    inhalt = 'Immobilie,Wohnung,Fläche (m²)\r\nHaus Müller,EG,50\r\n'.encode('cp1252')
    gelesen = _lesen(auth_client, 'wohnungen', inhalt)
    assert gelesen['kopf'] == ['Immobilie', 'Wohnung', 'Fläche (m²)']
    assert gelesen['beispiel'] == [['Haus Müller', 'EG', '50']]


def test_ohne_anmeldung_gesperrt(anon_client):
    assert anon_client.get('/api/import/arten').status_code == 401
    assert anon_client.post('/api/import/wohnungen/lesen').status_code == 401


# --- NK-230: Wie im Vorjahr füllen -------------------------------------------


def test_vorjahr_liefert_die_zeilen_ohne_betraege(haus):
    """Die Rechnungen des Vorjahres als Zeilen der Tabelle, Zeitraum ein Jahr
    weiter, Betrag, Nummer und Rechnungsdatum leer: nichts geschätzt."""
    from nebenkostenfix.models import (Apartment, CostCategory, CostInvoice, Heizungsanlage,
                                       Property, Provider, db)
    haus_id = Property.query.one().id
    eg = Apartment.query.one()
    art = {k.name: k for k in CostCategory.query.all()}
    stadt = Provider(name='Stadtkasse')
    anlage = Heizungsanlage(property_id=haus_id, name='Zentralheizung')
    db.session.add_all([stadt, anlage])
    db.session.flush()

    def rechnung(kostenart, von, bis, **kw):
        db.session.add(CostInvoice(property_id=haus_id, category_id=art[kostenart].id,
                                   amount=Decimal('100'), start_date=von, end_date=bis, **kw))
    rechnung('Grundsteuer', date(2024, 1, 1), date(2024, 12, 31), provider_id=stadt.id,
             invoice_number='GS-17', rechnungsdatum=date(2024, 1, 15))
    rechnung('Hauswart', date(2024, 2, 1), date(2024, 2, 29), apartment_id=eg.id)
    rechnung('Grundsteuer', date(2023, 1, 1), date(2023, 12, 31))   # vorvoriges Jahr
    rechnung('Grundsteuer', date(2025, 1, 1), date(2025, 12, 31))   # schon erfasst
    rechnung('Heizung', date(2024, 1, 1), date(2024, 12, 31), heizungsanlage_id=anlage.id,
             heizkostenart='brennstoff')
    db.session.commit()

    antwort = haus.get(f'/api/import/rechnungen/vorjahr?property_id={haus_id}&jahr=2025')
    assert antwort.status_code == 200
    daten = antwort.get_json()
    leer = {'betrag': '', 'rechnungsnummer': '', 'rechnungsdatum': ''}
    assert daten['zeilen'] == [
        {'kostenart': 'Grundsteuer', 'von': '01.01.2025', 'bis': '31.12.2025',
         'anbieter': 'Stadtkasse', 'wohnung': '', **leer},
        # Schaltjahr: der 29.02. wird zum 28.02.
        {'kostenart': 'Hauswart', 'von': '01.02.2025', 'bis': '28.02.2025',
         'anbieter': '', 'wohnung': 'EG links', **leer},
    ]
    # Heizkosten einer Anlage gehen nur über den Dialog (Anlage, CO2).
    assert daten['heizung'] == 1


def test_vorjahr_braucht_immobilie_und_jahr(haus):
    assert haus.get('/api/import/rechnungen/vorjahr?jahr=2025').status_code == 400
    assert haus.get('/api/import/rechnungen/vorjahr?property_id=1&jahr=x').status_code == 400
    assert haus.get('/api/import/rechnungen/vorjahr?property_id=999&jahr=2025').status_code == 404


def test_ein_jahr_weiter_haelt_das_monatsende():
    """NK-236: der 28.02. vor einem Schaltjahr wird zum 29.02., sonst fiele
    der letzte Tag des Zeitraums weg."""
    assert ti._ein_jahr_weiter(date(2027, 2, 28)) == date(2028, 2, 29)
    assert ti._ein_jahr_weiter(date(2028, 2, 29)) == date(2029, 2, 28)
    assert ti._ein_jahr_weiter(date(2027, 2, 15)) == date(2028, 2, 15)
    assert ti._ein_jahr_weiter(date(2027, 12, 31)) == date(2028, 12, 31)
